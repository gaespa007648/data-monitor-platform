"""Check exported files independently of Streamlit's browser widget."""

from io import BytesIO
from unittest.mock import patch
from zipfile import is_zipfile

import altair as alt
from openpyxl import load_workbook

from analysis_ui import apply_filters, render_analysis, render_chart
from export_utils import chart_png_bytes, excel_bytes, safe_filename_part


def main():
    dataset = {"fields": [
        {"name": "name", "data_type": "string"},
        {"name": "value", "data_type": "decimal"},
        {"name": "date", "data_type": "datetime"},
        {"name": "key", "data_type": "integer"},
        {"name": "active", "data_type": "boolean"},
    ]}
    rows = [
        {"name": "=1+1", "value": 1.25, "date": "2026-10-28T08:00:00+08:00",
         "key": 1234567890123456, "active": True},
        {"name": "ordinary", "value": 2.5, "date": "2026-10-28T09:00:00+08:00",
         "key": 2, "active": False},
    ]
    filtered = apply_filters(rows, {"value": ("decimal", (1, 2))})
    assert len(filtered) == 1
    payload = excel_bytes(dataset, filtered)
    assert is_zipfile(BytesIO(payload))
    workbook = load_workbook(BytesIO(payload), read_only=True)
    sheet = workbook.active
    assert len(list(sheet.rows)) == 2
    assert [cell.value for cell in sheet[1]] == ["name", "value", "date", "key", "active"]
    assert sheet["A2"].value == "=1+1" and sheet["A2"].data_type == "s"
    assert sheet["B2"].value == 1.25
    assert sheet["C2"].value.isoformat() == "2026-10-28T00:00:00"
    assert sheet["D2"].value == "1234567890123456"
    assert sheet["E2"].value is True
    workbook.close()

    empty_book = load_workbook(BytesIO(excel_bytes(dataset, [])), read_only=True)
    assert len(list(empty_book.active.rows)) == 1
    empty_book.close()
    assert safe_filename_part('iris/stock:2026') == "iris_stock_2026"

    chart = alt.Chart(alt.Data(values=[{"x": "a", "y": 1}])).mark_bar().encode(x="x:N", y="y:Q")
    assert chart_png_bytes(chart.to_json()).startswith(b"\x89PNG\r\n\x1a\n")

    dataset.update({"id": 7, "name": "iris/stocks"})
    with patch("analysis_ui.load_records", return_value=rows), patch("analysis_ui.render_statistics"):
        with patch("analysis_ui.render_filters", return_value={}), patch("analysis_ui.st.download_button") as button:
            render_analysis(dataset, None)
            assert button.call_args.kwargs["disabled"] is True
        with patch("analysis_ui.render_filters", return_value={"value": ("decimal", (1, 2))}), \
             patch("analysis_ui.st.download_button") as button:
            render_analysis(dataset, None)
            options = button.call_args.kwargs
            assert options["disabled"] is False
            assert options["file_name"] == "iris_stocks-filtered.xlsx"
            assert len(list(load_workbook(BytesIO(options["data"]), read_only=True).active.rows)) == 2

    with patch("analysis_ui.st.altair_chart"), patch("analysis_ui.st.download_button") as button:
        render_chart(chart, dataset, "value-histogram")
        options = button.call_args.kwargs
        assert options["file_name"] == "iris_stocks-value-histogram.png"
        assert options["mime"] == "image/png"
        assert options["data"].startswith(b"\x89PNG\r\n\x1a\n")
    print("Excel and PNG export checks passed")


if __name__ == "__main__":
    main()
